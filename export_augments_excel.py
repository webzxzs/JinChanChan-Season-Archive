"""Append or refresh the augment sheet in the existing five-tab workbook.

Optional research/export dependency: openpyxl. The Web viewer needs none.
Existing four worksheets, table styles, and formulas are retained.
"""
import csv
import os
from pathlib import Path

from augment_catalog import DEST, HEADERS, ROOT

XLSX = ROOT / "金铲铲全赛季研究资料库.xlsx"
SHEET = "海克斯强化"


def main():
    try:
        from openpyxl import load_workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.worksheet.table import Table, TableStyleInfo
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("更新Excel第5张表需安装可选依赖：pip install openpyxl") from exc

    if not XLSX.exists():
        raise FileNotFoundError(f"原四表 Excel 不存在: {XLSX}")
    with DEST.open(encoding="utf-8-sig", newline="") as stream:
        records = list(csv.reader(stream))
    if records[0] != HEADERS:
        raise ValueError("海克斯 CSV 列名不匹配")
    book = load_workbook(XLSX)
    required = ("棋子", "羁绊与效果", "棋盘可能效果", "赛季特色")
    if any(name not in book.sheetnames for name in required):
        raise ValueError("原四个工作表不全，拒绝修改已有研究工作簿")
    if SHEET in book.sheetnames:
        book.remove(book[SHEET])
    sheet = book.create_sheet(SHEET)
    for values in records:
        sheet.append(values)
    sheet.freeze_panes = "D2"
    sheet.sheet_view.zoomScale = 85
    widths = [38, 19, 17, 24, 85, 12, 24, 58, 22, 19, 17, 20, 35, 68, 62]
    for i, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(i)].width = width
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="273850")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 28
    for row in sheet.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=c.column in (5, 14))
        if row[5].value in ("1", "2", "3"):
            row[6].font = Font(color={"1": "6E8494", "2": "A17A23", "3": "8551B3"}[row[5].value])
    table = Table(displayName="AugmentCatalog", ref=f"A1:O{len(records)}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)
    tmp = XLSX.with_name(XLSX.stem + ".part.xlsx")
    try:
        book.save(tmp)
        check = load_workbook(tmp, read_only=True)
        assert all(name in check.sheetnames for name in required), "Original sheets missing"
        assert check[SHEET].max_row == len(records), "Augment row count mismatch"
        check.close()
        os.replace(tmp, XLSX)
    finally:
        if tmp.exists():
            tmp.unlink()
    print("Excel 海克斯强化 Sheet:", len(records) - 1, "条")


if __name__ == "__main__":
    main()
