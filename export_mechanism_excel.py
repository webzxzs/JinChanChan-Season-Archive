"""Refresh the sixth (classified) mechanism sheet without touching raw 5 tabs."""
import csv
import os

from mechanism_catalog import DEST, ROOT

WORKBOOK = ROOT / "金铲铲全赛季研究资料库.xlsx"
SHEET = "赛季机制图鉴"


def main():
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    with DEST.open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.reader(stream))
    if not rows or len(rows[0]) != 16:
        raise ValueError("机制图鉴 CSV 结构异常，拒绝覆盖 Excel")
    book = load_workbook(WORKBOOK)
    base = ("棋子", "羁绊与效果", "棋盘可能效果", "赛季特色", "海克斯强化")
    if any(k not in book.sheetnames for k in base):
        raise ValueError("工作簿五张基础表缺失")
    if SHEET in book.sheetnames:
        book.remove(book[SHEET])
    sheet = book.create_sheet(SHEET)
    for row in rows:
        sheet.append(["'" + x if x.startswith(("=", "+", "@")) else x for x in row])
    sheet.freeze_panes = "D2"
    sheet.sheet_view.zoomScale = 85
    widths = [38, 19, 22, 28, 72, 42, 50, 26, 65, 62, 43, 26, 35, 23, 19, 68]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="273850")
        cell.font = Font(bold=True, color="FFFFFF")
    for row in sheet.iter_rows(min_row=2):
        row[4].alignment = Alignment(wrap_text=True, vertical="top")
        row[15].alignment = Alignment(wrap_text=True, vertical="top")
    table = Table(displayName="MechanismCatalog", ref=f"A1:P{len(rows)}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)
    out = WORKBOOK.with_name(WORKBOOK.stem + ".classified.xlsx")
    try:
        book.save(out)
        check = load_workbook(out, read_only=True)
        if check[SHEET].max_row != len(rows) or any(n not in check.sheetnames for n in base):
            raise ValueError("生成工作簿行数或原表数量不正确")
        check.close()
        os.replace(out, WORKBOOK)
    finally:
        if out.exists():
            out.unlink()
    print(f"Excel 已更新「{SHEET}」: {len(rows)-1} 条")


if __name__ == "__main__":
    main()
